import sys, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

main_f, pm_f, snap_f, out_f = sys.argv[1:5]
F = 'Arial'


def g(r, i):
    return r[i] if len(r) > i else None


def num(v):
    try:
        f = float(v)
        return f if f > 0 else None
    except (TypeError, ValueError):
        return None


def key(s):
    return str(s).strip().upper()


# ── sources ──────────────────────────────────────────────────────────────
main = {}
wb = openpyxl.load_workbook(main_f, read_only=True, data_only=True)
for r in wb['Main'].iter_rows(min_row=4, values_only=True):
    sku = g(r, 1)
    if not isinstance(sku, str) or not sku.strip() or g(r, 2) is None and num(g(r, 17)) is None:
        continue
    main.setdefault(key(sku), {'sku': sku.strip(), 'unit': num(g(r, 17)), 'pq': num(g(r, 13)), 'plb': num(g(r, 25)), 'desc': g(r, 2)})

wb2 = openpyxl.load_workbook(pm_f, read_only=True, data_only=True)
pm = {}
for sh in ['Riverside Copy', 'Units per case pallet ']:
    for r in wb2[sh].iter_rows(min_row=2, values_only=True):
        sku = g(r, 0)
        if sku and key(sku) not in pm:
            pm[key(sku)] = {'sku': str(sku).strip(), 'pq': num(g(r, 3)), 'plb': num(g(r, 6)), 'desc': g(r, 1), 'sheet': sh.strip()}
acc = {}
for r in wb2['Accesories II'].iter_rows(min_row=5, values_only=True):
    sku = g(r, 0)
    if sku and str(sku) != 'EXAMPLE-SKU':
        acc[key(sku)] = {'sku': str(sku).strip(), 'pq': num(g(r, 4)), 'plb': num(g(r, 7)), 'desc': g(r, 1)}

snap = json.load(open(snap_f))
names = {str(i['item']): i for i in snap['items']}
onhand = {}
for o in snap['onHand']:
    it = names.get(str(o['item']), {})
    if it.get('sku'):
        onhand[key(it['sku'])] = (it['sku'], it.get('descr'), o['onhand'])

# ── resolve one weight per SKU (priority order documented on Notes) ──────
rows = []      # (sku, desc, kind, value-or-ref, pq, plb, source)
done = set()
for k in sorted(set(main) | set(pm) | set(acc)):
    m, p, a = main.get(k), pm.get(k), acc.get(k)
    desc = (m or p or a)['desc']
    sku = (m or p or a)['sku']
    if m and m['unit']:
        rows.append((sku, desc, 'unit', m['unit'], m['pq'], m['plb'], 'Products Main: individual weight'))
    elif m and m['plb'] and m['pq']:
        rows.append((sku, desc, 'pallet', None, m['pq'], m['plb'], 'Products Main: pallet weight ÷ pallet qty'))
    elif p and p['plb'] and p['pq']:
        rows.append((sku, desc, 'pallet', None, p['pq'], p['plb'], 'Product Matrix (' + p['sheet'] + '): pallet weight ÷ pallet qty'))
    elif a and a['plb'] and a['pq']:
        rows.append((sku, desc, 'pallet', None, a['pq'], a['plb'], 'Product Matrix (Accesories II): pallet weight ÷ pallet qty'))
    else:
        continue
    done.add(k)
# variants on hand at Riverside inherit the base SKU (YSN201-AG → YSN201)
for k, (sku, desc, q) in sorted(onhand.items()):
    if k in done or '-' not in k:
        continue
    base = k.rsplit('-', 1)[0]
    if base in done:
        rows.append((sku, desc, 'variant', base, None, None, 'From base SKU ' + base))
        done.add(k)
rows.sort(key=lambda r: key(r[0]))

# ── workbook ─────────────────────────────────────────────────────────────
out = openpyxl.Workbook()
ws = out.active
ws.title = 'Unit weights'
hdr = ['SKU', 'Unit weight (lb)', 'Source', 'Description', 'Pallet qty', 'Pallet weight (lb)', 'Riverside on hand (10/06)']
ws.append(hdr)
blue = Font(name=F, color='0000FF')
black = Font(name=F)
for c in ws[1]:
    c.font = Font(name=F, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', start_color='1F3864')
    c.alignment = Alignment(wrap_text=True, vertical='center')
row_of = {}
for i, (sku, desc, kind, val, pq, plb, src) in enumerate(rows, start=2):
    row_of[key(sku)] = i
for i, (sku, desc, kind, val, pq, plb, src) in enumerate(rows, start=2):
    oh = onhand.get(key(sku))
    ws.cell(i, 1, sku).font = black
    ws.cell(i, 3, src).font = black
    ws.cell(i, 4, desc).font = black
    if pq:
        ws.cell(i, 5, pq).font = blue
    if plb:
        ws.cell(i, 6, plb).font = blue
    if oh:
        ws.cell(i, 7, oh[2]).font = black
    c = ws.cell(i, 2)
    if kind == 'unit':
        c.value, c.font = val, blue
    elif kind == 'pallet':
        c.value, c.font = '=ROUND(F%d/E%d,2)' % (i, i), black
    else:
        c.value, c.font = '=B%d' % row_of[val], black
    c.number_format = '0.00'
    ws.cell(i, 5).number_format = '#,##0'
    ws.cell(i, 6).number_format = '#,##0'
    ws.cell(i, 7).number_format = '#,##0'
last = len(rows) + 1
for col, w in zip('ABCDEFG', [24, 14, 44, 60, 11, 14, 14]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = 'A1:G%d' % last
ws['B1'].comment = Comment('Blue = typed from your files. Black = formula (pallet weight ÷ pallet qty, or the base SKU for variants).', 'Claude')

# Missing sheet: Riverside on-hand SKUs with no weight anywhere
ms = out.create_sheet('Missing (Riverside)')
ms['A1'] = 'Riverside SKUs on hand with no weight in either file. Fill in the yellow column, then copy the rows into Unit weights.'
ms['A1'].font = Font(name=F, italic=True)
ms.append([])
ms.append(['SKU', 'Description', 'Riverside on hand (10/06)', 'Unit weight (lb)', 'Notes'])
for c in ms[3]:
    c.font = Font(name=F, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', start_color='1F3864')
ms.append(['EXAMPLE-SKU', 'Example: weigh one unit on the floor scale', 1200, 2.5, 'Example row: delete before upload'])
for c in ms[4]:
    c.font = Font(name=F, italic=True, color='808080')
yellow = PatternFill('solid', start_color='FFFF00')
miss = sorted([v for k, v in onhand.items() if k not in done], key=lambda v: -(v[2] or 0))
for sku, desc, q in miss:
    ms.append([sku, desc, q, None, None])
    r = ms.max_row
    for c in ms[r]:
        c.font = black
    ms.cell(r, 4).fill = yellow
    ms.cell(r, 3).number_format = '#,##0'
for col, w in zip('ABCDE', [26, 60, 14, 14, 30]):
    ms.column_dimensions[col].width = w
ms.freeze_panes = 'A4'

# Notes / coverage
ns = out.create_sheet('Notes')
tot_rows = len(snap['onHand'])
lines = [
    ('Combined unit weights for the Move Portal', True),
    ('Built 2026-10-19 from "Flame King Products Main (1).xlsx" (sheet Main) and "Product Matrix .xlsx" (sheets Riverside Copy, Units per case pallet, Accesories II).', False),
    ('', False),
    ('Which number wins (first match per SKU):', True),
    ('1. Products Main, INDIVIDUAL WEIGHT (LB) column.', False),
    ('2. Products Main, PALLET WEIGHT (LB) ÷ PALLET QUANTITY.', False),
    ('3. Product Matrix, Riverside Copy (then Units per case pallet): Pallet weight ÷ Pallet Qty.', False),
    ('4. Product Matrix, Accesories II: Pallet Weight ÷ Pallet Qty.', False),
    ('5. Colour/label variants on hand at Riverside (e.g. YSN201-AG, -BR, -UH) use their base SKU (Jack, 2026-10-19).', False),
    ('', False),
    ('Assumptions to check:', True),
    ('• A pallet weight may include the wooden pallet. If so, a pallet-derived unit weight runs slightly high (about 45 lb ÷ pallet qty per unit). The portal already adds its own pallet tare (Settings palletTare, 45 lb).', False),
    ('• YSN201 = 18 lb (Products Main). Capacity Builder said 19.8 lb; the real BOL used 18 lb.', False),
    ('', False),
    ('Coverage (Riverside on hand, snapshot 2026-10-06):', True),
    ('SKUs on hand', False), ('SKUs with a weight', False), ('SKUs missing', False), ('Pieces covered', False),
]
for t, bold in lines:
    ns.append([t])
    ns.cell(ns.max_row, 1).font = Font(name=F, bold=bold)
r0 = ns.max_row - 3
ns.cell(r0, 2, tot_rows)
ns.cell(r0 + 1, 2, "=%d-B%d" % (tot_rows, r0 + 2))
ns.cell(r0 + 2, 2, "=COUNTA('Missing (Riverside)'!A5:A1000)")
ns.cell(r0 + 3, 2, "=SUM('Unit weights'!G2:G%d)/(SUM('Unit weights'!G2:G%d)+SUM('Missing (Riverside)'!C5:C1000))" % (last, last))
ns.cell(r0 + 3, 2).number_format = '0.0%'
for rr in range(r0, r0 + 4):
    ns.cell(rr, 2).font = black
ns.column_dimensions['A'].width = 110
ns.column_dimensions['B'].width = 12
out.move_sheet('Notes', offset=-2)
out.active = 1
out.save(out_f)
print('rows', len(rows), 'missing', len(miss), 'out', out_f)
