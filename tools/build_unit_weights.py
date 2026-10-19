"""Build the combined unit-weights workbook for the Move Portal.

usage: python build_unit_weights.py main.xlsx matrix.xlsx onhand.json out.xlsx

onhand.json is [{sku, descr, onhand}] (Riverside on hand, live SuiteQL).
Column B of the Unit weights sheet holds VALUES (the portal upload reads cached
values only); column I holds the live formula for audit.
"""
import sys
import json
import re
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

if len(sys.argv) != 5:
    sys.exit('usage: build_unit_weights.py main.xlsx matrix.xlsx onhand.json out.xlsx')
main_f, pm_f, oh_f, out_f = sys.argv[1:5]
F = 'Arial'
TARE = 35
CYL_SECTION = ('CYLINDER', 'TANK', 'ASME', 'FORKLIFT')
NOT_CYL = ('VALVE', 'GAUGE', 'PART', 'ACCESSOR', 'KIT', 'REFILLABLE')


def sec_type(name):
    u = name.upper()
    return 'CYL' if any(w in u for w in CYL_SECTION) and not any(w in u for w in NOT_CYL) else 'PCS'


CYL_WORDS = re.compile(r'\b(cylinders?|tanks?)\b', re.I)


def g(r, i):
    return r[i] if len(r) > i else None


def num(v):
    try:
        f = float(v)
        return f if f > 0 else None
    except (TypeError, ValueError):
        return None


def key(s):
    return str(s).strip().upper()


def rnd2(x):
    return float(Decimal(repr(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def mtype(sku, desc):
    t = sku + ' ' + (desc or '')
    return 'CYL' if CYL_WORDS.search(t) and not any(w in t.upper() for w in NOT_CYL) else 'PCS'


# ── sources ──────────────────────────────────────────────────────────────
main = {}
sections = {}   # section name -> Type
section = None
wb = openpyxl.load_workbook(main_f, read_only=True, data_only=True)
for r in wb['Main'].iter_rows(min_row=4, values_only=True):
    sku = g(r, 1)
    if not isinstance(sku, str) or not sku.strip():
        continue
    if (g(r, 2) is None and g(r, 0) is None and num(g(r, 13)) is None and num(g(r, 17)) is None
            and num(g(r, 21)) is None and num(g(r, 25)) is None
            and not any(isinstance(v, (int, float)) for v in r)):
        section = sku.strip()
        sections[section] = sec_type(section)
        continue
    main.setdefault(key(sku), {'sku': sku.strip(), 'unit': num(g(r, 17)), 'pq': num(g(r, 13)), 'plb': num(g(r, 25)),
                               'desc': g(r, 2), 'type': sections.get(section, 'PCS')})

wb2 = openpyxl.load_workbook(pm_f, read_only=True, data_only=True)
pm = {}
for sh in ['Riverside Copy', 'Units per case pallet ']:
    for r in wb2[sh].iter_rows(min_row=2, values_only=True):
        sku = g(r, 0)
        if sku and key(sku) not in pm:
            pm[key(sku)] = {'sku': str(sku).strip(), 'pq': num(g(r, 3)), 'plb': num(g(r, 6)), 'desc': g(r, 1), 'sheet': sh.strip()}
acc = {}
for r in wb2['Accesories II'].iter_rows(min_row=5, values_only=True):
    sku = g(r, 0)
    if sku and str(sku) != 'EXAMPLE-SKU':
        acc[key(sku)] = {'sku': str(sku).strip(), 'pq': num(g(r, 4)), 'plb': num(g(r, 7)), 'desc': g(r, 1)}

onhand = {}
for o in json.load(open(oh_f)):
    k = key(o['sku'])
    prev = onhand.get(k)
    onhand[k] = (o['sku'], o.get('descr'), o['onhand'] + (prev[2] if prev else 0))

# ── resolve one weight per SKU (priority order documented on Notes) ──────
rows = []
done = {}
for k in sorted(set(main) | set(pm) | set(acc)):
    m, p, a = main.get(k), pm.get(k), acc.get(k)
    desc = (m or p or a)['desc']
    sku = (m or p or a)['sku']
    row = {'sku': sku, 'desc': desc, 'type': m['type'] if m else mtype(sku, desc), 'matrix_only': not m}
    if m and m['unit']:
        row.update(kind='unit', val=m['unit'], pq=m['pq'], plb=m['plb'], src='Products Main: individual weight')
    else:
        cands = [(m, 'Products Main: pallet weight less tare, / pallet qty'),
                 (p, 'Product Matrix (%s): pallet weight less tare, / pallet qty' % (p or {}).get('sheet')),
                 (a, 'Product Matrix (Accesories II): pallet weight less tare, / pallet qty')]
        for src, label in cands:
            if src and src['plb'] and src['pq']:
                row.update(kind='pallet', pq=src['pq'], plb=src['plb'], src=label,
                           val=rnd2((src['plb'] - TARE) / src['pq']))
                break
        else:
            continue
    rows.append(row)
    done[k] = row
# variants on hand at Riverside inherit the base SKU (YSN201-AG -> YSN201)
for k, (sku, desc, q) in sorted(onhand.items()):
    if k in done or '-' not in k:
        continue
    base = k.rsplit('-', 1)[0]
    if base in done:
        b = done[base]
        row = {'sku': sku, 'desc': desc, 'type': b['type'], 'kind': 'variant', 'base': base, 'val': b['val'],
               'src': 'From base SKU ' + b['sku'], 'pq': None, 'plb': None, 'matrix_only': False}
        rows.append(row)
        done[k] = row
rows.sort(key=lambda r: key(r['sku']))

# ── workbook ─────────────────────────────────────────────────────────────
out = openpyxl.Workbook()
ws = out.active
ws.title = 'Unit weights'
ws.append(['SKU', 'Unit weight (lb)', 'Type', 'Source', 'Description', 'Pallet qty', 'Pallet weight (lb)',
           'Riverside on hand', 'Calc (formula)'])
blue = Font(name=F, color='0000FF')
black = Font(name=F)
for c in ws[1]:
    c.font = Font(name=F, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', start_color='1F3864')
    c.alignment = Alignment(wrap_text=True, vertical='center')
row_of = {key(r['sku']): i for i, r in enumerate(rows, start=2)}
for i, r in enumerate(rows, start=2):
    oh = onhand.get(key(r['sku']))
    ws.cell(i, 1, r['sku']).font = black
    ws.cell(i, 3, r['type']).font = black
    ws.cell(i, 4, r['src']).font = black
    ws.cell(i, 5, r['desc']).font = black
    if r['pq']:
        ws.cell(i, 6, r['pq']).font = blue
    if r['plb']:
        ws.cell(i, 7, r['plb']).font = blue
    if oh:
        ws.cell(i, 8, oh[2]).font = black
    c = ws.cell(i, 2, r['val'])
    c.number_format = '0.00'
    if r['kind'] == 'unit':
        c.font = blue
    else:
        c.font = black
        f = ('=ROUND((G%d-Notes!$B$3)/F%d,2)' % (i, i)) if r['kind'] == 'pallet' else '=B%d' % row_of[r['base']]
        ws.cell(i, 9, f).font = black
    for col in (6, 7, 8):
        ws.cell(i, col).number_format = '#,##0'
last = len(rows) + 1
for col, w in zip('ABCDEFGHI', [24, 14, 8, 52, 60, 11, 14, 14, 30]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
ws.auto_filter.ref = 'A1:I%d' % last
ws['B1'].comment = Comment('Blue = typed from your files. Black = computed (pallet weight less tare, divided by pallet qty; or the '
                           'base SKU for variants). Column I shows the live formula for audit.', 'Claude')

# Missing sheet: Riverside on-hand SKUs with no weight anywhere
ms = out.create_sheet('Missing (Riverside)')
ms['A1'] = ('Riverside SKUs on hand with no weight in either file. Fill in the yellow columns (weight, type CYL or PCS), '
            'then copy the rows into Unit weights.')
ms['A1'].font = Font(name=F, italic=True)
ms.append([])
ms.append(['SKU', 'Description', 'Riverside on hand', 'Unit weight (lb)', 'Type', 'Notes'])
for c in ms[3]:
    c.font = Font(name=F, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', start_color='1F3864')
ms.append(['EXAMPLE-SKU', 'Example: weigh one unit on the floor scale', 1200, 2.5, 'PCS', 'Example row: delete before upload'])
for c in ms[4]:
    c.font = Font(name=F, italic=True, color='808080')
yellow = PatternFill('solid', start_color='FFFF00')
miss = sorted([v for k, v in onhand.items() if k not in done], key=lambda v: -(v[2] or 0))
for sku, desc, q in miss:
    ms.append([sku, desc, q, None, None, None])
    r = ms.max_row
    for c in ms[r]:
        c.font = black
    ms.cell(r, 4).fill = yellow
    ms.cell(r, 5).fill = yellow
    ms.cell(r, 3).number_format = '#,##0'
for col, w in zip('ABCDEF', [26, 60, 14, 14, 8, 30]):
    ms.column_dimensions[col].width = w
ms.freeze_panes = 'A4'

# Notes / coverage (values; the tare in B3 is the one live cell, used by the Calc column)
ns = out.create_sheet('Notes')
tot_skus = len(onhand)
tot_pcs = sum(v[2] for v in onhand.values())
cov_pcs = sum(v[2] for k, v in onhand.items() if k in done)
ns.append(['Combined unit weights for the Move Portal'])
ns.append(['Built 2026-10-19 from "Flame King Products Main (1).xlsx" (sheet Main) and "Product Matrix .xlsx" (sheets Riverside Copy, '
           'Units per case pallet, Accesories II), with Riverside on hand read live from NetSuite on 2026-10-19.'])
ns.append(['Pallet tare (lb)', TARE])
ns.append(['Column B on Unit weights holds values for the portal upload. Changing the tare here updates the Calc column only; '
           'rebuild with tools/build_unit_weights.py to change column B.'])
bold_rows = []
for t, bold in [
    ('', False),
    ('Which number wins (first match per SKU):', True),
    ('1. Products Main, INDIVIDUAL WEIGHT (LB) column.', False),
    ('2. Products Main, (PALLET WEIGHT LB - tare) / PALLET QUANTITY.', False),
    ('3. Product Matrix, Riverside Copy (then Units per case pallet): (Pallet weight - tare) / Pallet Qty.', False),
    ('4. Product Matrix, Accesories II: (Pallet Weight - tare) / Pallet Qty.', False),
    ('5. Colour/label variants on hand at Riverside (e.g. YSN201-AG, -BR, -UH) use their base SKU (Jack, 2026-10-19).', False),
    ('', False),
    ('Type (BOL): CYL for Products Main sections whose name has CYLINDER, TANK, ASME or FORKLIFT and none of VALVE, GAUGE, PART, ACCESSOR, KIT, REFILLABLE; everything else '
     'PCS. Matrix-only SKUs are PCS unless the SKU or description says cylinder/tank (same exclusions). Variants copy the base SKU Type. '
     'Review the Type column.', False),
    ('', False),
    ('Assumptions to check:', True),
    ('- YSN201 = 18 lb (Products Main). Capacity Builder said 19.8 lb; the real BOL used 18 lb.', False),
    ('', False),
    ('Coverage (Riverside on hand, live 2026-10-19):', True),
]:
    ns.append([t])
    if bold:
        bold_rows.append(ns.max_row)
for lab, val, fmt in (('SKUs on hand', tot_skus, '#,##0'), ('SKUs with a weight', tot_skus - len(miss), '#,##0'),
                      ('SKUs missing', len(miss), '#,##0'), ('Pieces covered', cov_pcs / tot_pcs, '0.0%')):
    ns.append([lab, val])
    ns.cell(ns.max_row, 2).number_format = fmt
for rr in range(1, ns.max_row + 1):
    for cc in (1, 2):
        ns.cell(rr, cc).font = black
ns['A1'].font = Font(name=F, bold=True, size=12)
ns['A3'].font = Font(name=F, bold=True)
ns['B3'].font = blue
for rr in bold_rows:
    ns.cell(rr, 1).font = Font(name=F, bold=True)
ns.column_dimensions['A'].width = 110
ns.column_dimensions['B'].width = 12
# Notes stays last: the portal upload picks the sheet named Unit weights, and it is the first sheet too
out.active = 0
out.save(out_f)
print('rows', len(rows), 'missing', len(miss), 'out', out_f)
print('sections', json.dumps(sections, indent=1))
print('coverage skus %d/%d pieces %.1f%%' % (tot_skus - len(miss), tot_skus, 100.0 * cov_pcs / tot_pcs))
print('matrix-only CYL:', [r['sku'] for r in rows if r['matrix_only'] and r['type'] == 'CYL'])
